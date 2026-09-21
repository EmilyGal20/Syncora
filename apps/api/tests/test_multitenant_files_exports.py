import io
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from fastapi import HTTPException, UploadFile
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine
from starlette.datastructures import Headers

from app.database import Base
from app.models import BandShow, Dashboard, Organization, Permission, Setlist, Song, StoredFile, Task, User
from app.routers import band, platform
from app.schemas import BandResourceInput, FileAttachmentInput, FileRename
from app.storage import LocalStorageAdapter


@pytest.fixture
async def database(tmp_path):
    engine = create_async_engine(f"sqlite+aiosqlite:///{tmp_path / 'test.db'}")
    async with engine.begin() as connection:
        await connection.run_sync(Base.metadata.create_all)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    async with factory() as session:
        yield session
    await engine.dispose()


def upload(name: str, content_type: str, content: bytes) -> UploadFile:
    return UploadFile(filename=name, file=io.BytesIO(content), headers=Headers({"content-type": content_type}))


def test_local_storage_rejects_path_traversal(tmp_path):
    storage = LocalStorageAdapter(str(tmp_path / "files"))
    with pytest.raises(ValueError):
        storage._path("../outside.txt")


@pytest.mark.asyncio
async def test_upload_validates_type_size_and_generates_tenant_key(tmp_path, monkeypatch):
    storage = LocalStorageAdapter(str(tmp_path / "files"))
    item = await platform.save_upload(
        upload("stage plot.pdf", "application/pdf", b"%PDF-test"), "band-a", "user-a", "show", "show-a", storage
    )
    assert item.storage_key.startswith("band-a/")
    assert item.original_filename == "stage plot.pdf"
    with pytest.raises(HTTPException) as invalid:
        await platform.save_upload(
            upload("run.exe", "application/octet-stream", b"MZ"), "band-a", "user-a", "workspace", None, storage
        )
    assert invalid.value.status_code == 415
    monkeypatch.setattr(platform.settings, "upload_max_bytes", 3)
    with pytest.raises(HTTPException) as oversized:
        await platform.save_upload(
            upload("notes.txt", "text/plain", b"long"), "band-a", "user-a", "workspace", None, storage
        )
    assert oversized.value.status_code == 413
    monkeypatch.setattr(platform.settings, "upload_max_bytes", 100)
    with pytest.raises(HTTPException) as disguised:
        await platform.save_upload(
            upload("fake.pdf", "application/pdf", b"not a pdf"),
            "band-a",
            "user-a",
            "workspace",
            None,
            storage,
        )
    assert disguised.value.status_code == 415


@pytest.mark.asyncio
async def test_file_queries_and_downloads_are_tenant_qualified(database, tmp_path):
    org_a = Organization(name="Band Alpha", slug="alpha")
    org_b = Organization(name="Band Beta", slug="beta")
    database.add_all([org_a, org_b])
    await database.flush()
    user_a = User(organization_id=org_a.id, email="a@example.test", full_name="Alpha", password_hash="x")
    user_b = User(organization_id=org_b.id, email="b@example.test", full_name="Beta", password_hash="x")
    database.add_all([user_a, user_b])
    await database.flush()
    storage = LocalStorageAdapter(str(tmp_path / "files"))
    await storage.put(f"{org_b.id}/beta.pdf", b"beta secret")
    beta_file = StoredFile(
        organization_id=org_b.id,
        uploaded_by=user_b.id,
        original_filename="beta.pdf",
        display_name="beta.pdf",
        storage_key=f"{org_b.id}/beta.pdf",
        content_type="application/pdf",
        size=11,
        checksum="a" * 64,
    )
    database.add(beta_file)
    await database.commit()
    assert await platform.list_files(user=user_a, db=database) == []
    with pytest.raises(HTTPException) as denied:
        await platform.download_file(beta_file.id, user=user_a, db=database, storage=storage)
    assert denied.value.status_code == 404


async def response_bytes(response) -> bytes:
    return b"".join([part async for part in response.body_iterator])


@pytest.mark.asyncio
async def test_pdf_and_xlsx_are_valid_and_exclude_other_tenant(database, monkeypatch):
    monkeypatch.setattr(platform, "record_audit", AsyncMock())
    org_a = Organization(name="Band Alpha", slug="alpha")
    org_b = Organization(name="Band Beta", slug="beta")
    database.add_all([org_a, org_b])
    await database.flush()
    user = User(organization_id=org_a.id, email="member@alpha.test", full_name="Alpha Member", password_hash="x")
    database.add(user)
    await database.flush()
    database.add_all(
        [
            Dashboard(organization_id=org_a.id, name="Alpha dashboard", is_default=True),
            Task(organization_id=org_a.id, title="Alpha task", creator_id=user.id, assignee_id=user.id),
            Task(organization_id=org_b.id, title="Beta secret", creator_id=user.id, assignee_id=user.id),
        ]
    )
    await database.commit()
    user._workspace = org_a
    permission = Permission(code="tasks.view", group="Tasks")
    user._access_grants = [SimpleNamespace(permission=permission, principal_type="user", effect="allow", scope="OWN")]
    pdf = await response_bytes(await platform.export_pdf(user=user, db=database))
    assert pdf.startswith(b"%PDF")
    assert b"Beta secret" not in pdf
    xlsx = await response_bytes(await platform.export_xlsx(user=user, db=database))
    workbook = load_workbook(io.BytesIO(xlsx))
    assert workbook.sheetnames == ["Dashboard Summary", "Tasks"]
    values = [cell.value for cell in workbook["Tasks"]["A"]]
    assert "Alpha task" in values
    assert "Beta secret" not in values


def test_organization_admin_is_not_platform_admin():
    ordinary = SimpleNamespace(is_platform_admin=False)
    with pytest.raises(HTTPException) as denied:
        platform.require_platform(ordinary)
    assert denied.value.status_code == 403


@pytest.mark.asyncio
async def test_band_resource_ids_cannot_cross_tenant(database):
    org_a = Organization(name="Band Alpha", slug="alpha")
    org_b = Organization(name="Band Beta", slug="beta")
    database.add_all([org_a, org_b])
    await database.flush()
    alpha = User(organization_id=org_a.id, email="alpha@band.test", full_name="Alpha", password_hash="x")
    database.add(alpha)
    beta_show = BandShow(
        organization_id=org_b.id, title="Beta show", venue="Secret", starts_at=platform.datetime.now(platform.UTC)
    )
    beta_song = Song(organization_id=org_b.id, title="Beta song")
    beta_setlist = Setlist(organization_id=org_b.id, title="Beta setlist")
    database.add_all([beta_show, beta_song, beta_setlist])
    await database.commit()
    alpha._access_grants = [
        SimpleNamespace(
            permission=Permission(code=code, group="Band"), principal_type="user", effect="allow", scope="ORGANIZATION"
        )
        for code in ["shows.view", "songs.view", "setlists.view"]
    ]
    for kind, item_id in [("shows", beta_show.id), ("songs", beta_song.id)]:
        with pytest.raises(HTTPException) as denied:
            await band.resource(kind, item_id, user=alpha, db=database)
        assert denied.value.status_code == 404
    with pytest.raises(HTTPException) as setlist_denied:
        await band.setlist_detail(beta_setlist.id, user=alpha, db=database)
    assert setlist_denied.value.status_code == 404


@pytest.mark.asyncio
async def test_file_rename_attach_conflict_detach_and_delete(database, tmp_path, monkeypatch):
    monkeypatch.setattr(platform, "record_audit", AsyncMock())
    org = Organization(name="Band Alpha", slug="alpha-files")
    database.add(org)
    await database.flush()
    user = User(organization_id=org.id, email="files@alpha.test", full_name="Files Admin", password_hash="x")
    song = Song(organization_id=org.id, title="Opening Song")
    database.add_all([user, song])
    await database.flush()
    storage = LocalStorageAdapter(str(tmp_path / "files"))
    key = f"{org.id}/notes.txt"
    await storage.put(key, b"stage notes")
    item = StoredFile(
        organization_id=org.id,
        uploaded_by=user.id,
        original_filename="notes.txt",
        display_name="notes.txt",
        storage_key=key,
        content_type="text/plain",
        size=11,
        checksum="b" * 64,
    )
    database.add(item)
    await database.commit()

    renamed = await platform.rename_file(item.id, FileRename(display_name="Stage notes.txt"), user=user, db=database)
    assert renamed["display_name"] == "Stage notes.txt"
    body = FileAttachmentInput(context_type="song", context_id=song.id)
    link = await platform.attach_file(item.id, body, user=user, db=database)
    with pytest.raises(HTTPException) as duplicate:
        await platform.attach_file(item.id, body, user=user, db=database)
    assert duplicate.value.status_code == 409
    with pytest.raises(HTTPException) as referenced:
        await platform.delete_file(item.id, user=user, db=database, storage=storage)
    assert referenced.value.status_code == 409
    await platform.detach_file(item.id, link["id"], user=user, db=database)
    response = await platform.delete_file(item.id, user=user, db=database, storage=storage)
    assert response.status_code == 204


@pytest.mark.asyncio
async def test_show_crud_is_tenant_qualified_and_module_gated(database, monkeypatch):
    monkeypatch.setattr(band, "record_audit", AsyncMock())
    alpha_org = Organization(name="Alpha CRUD", slug="alpha-crud", enabled_modules=["shows"])
    beta_org = Organization(name="Beta CRUD", slug="beta-crud", enabled_modules=["shows"])
    database.add_all([alpha_org, beta_org])
    await database.flush()
    alpha = User(organization_id=alpha_org.id, email="crud@alpha.test", full_name="Alpha Admin", password_hash="x")
    beta = User(organization_id=beta_org.id, email="crud@beta.test", full_name="Beta Admin", password_hash="x")
    database.add_all([alpha, beta])
    await database.flush()
    permission = Permission(code="shows.manage", group="Band")
    grant = SimpleNamespace(permission=permission, principal_type="user", effect="allow", scope="ORGANIZATION")
    alpha._access_grants = [grant]
    beta._access_grants = [grant]
    alpha._workspace = alpha_org
    beta._workspace = beta_org
    starts = platform.datetime.now(platform.UTC)
    created = await band.create_resource(
        "shows", BandResourceInput(title="Alpha Show", venue="Club", starts_at=starts), user=alpha, db=database
    )
    updated = await band.update_resource(
        "shows",
        created.id,
        BandResourceInput(title="Alpha Show Updated", venue="Arena", starts_at=starts),
        user=alpha,
        db=database,
    )
    assert updated.title == "Alpha Show Updated"
    with pytest.raises(HTTPException) as cross_tenant:
        await band.delete_resource("shows", created.id, user=beta, db=database)
    assert cross_tenant.value.status_code == 404
    await band.delete_resource("shows", created.id, user=alpha, db=database)
    assert not await database.get(BandShow, created.id)
    alpha_org.enabled_modules = []
    with pytest.raises(HTTPException) as disabled:
        await band.resources("shows", user=alpha, db=database)
    assert disabled.value.status_code == 404
